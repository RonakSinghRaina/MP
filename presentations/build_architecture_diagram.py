"""Schematic of the hybrid model: how it builds on a basic U-Net, step by step.

Four slides on the lab template. Every box, arrow and label is a plain shape so
the whole diagram stays editable/resizable after import into Google Slides.
All sizes are measured from the code (src/hybrid_rfi_package/hybrid_model.py,
base 8, 512x512 input): 8->16->32->64 channels, bottleneck 128 x 32 x 32,
593,842 parameters.

    ~/torch-env/bin/python presentations/build_architecture_diagram.py
"""
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
TEMPLATE = os.path.join(ROOT, "Weekly_Report_Week2_DiL_Lab_v2.pptx")
OUT = os.path.join(ROOT, "Hybrid_Model_Schematic.pptx")
NAME = "Our model"          # change here once the new name is chosen

RED = RGBColor(0xE3, 0x18, 0x37)
PALE_RED = RGBColor(0xFB, 0xE3, 0xE7)
INK = RGBColor(0x1E, 0x1E, 0x1E)
NAVY = RGBColor(0x0F, 0x20, 0x43)
MIDNAVY = RGBColor(0x2E, 0x4A, 0x7D)
GREY = RGBColor(0x5A, 0x61, 0x70)
LIGHT = RGBColor(0xDD, 0xE1, 0xE8)
PANEL = RGBColor(0xF6, 0xF7, 0xF9)
EDGE = RGBColor(0xD9, 0xDD, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(TEMPLATE)
for sld in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)
LAY_OBJ = prs.slide_layouts[1]


# ------------------------------------------------------------------ helpers
def new_slide(title):
    s = prs.slides.add_slide(LAY_OBJ)
    tf = s.placeholders[0].text_frame
    tf.clear()
    r = tf.paragraphs[0].add_run()
    r.text = "➢ " + title
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Calibri", Pt(26), True, RED
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0 and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)
    return s


def text(sh, x, y, w, h, paras, size=12, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=2):
    box = sh.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        for run in ([(para, False)] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = run[0]
            r.font.name, r.font.size, r.font.bold = "Calibri", Pt(size), run[1]
            r.font.color.rgb = run[2] if len(run) > 2 else color
    return box


def box(sh, x, y, w, h, fill, lines, size=11, color=WHITE, line=None,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, bold_first=True):
    b = sh.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid()
    b.fill.fore_color.rgb = fill
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb = line
        b.line.width = Pt(1)
    b.shadow.inherit = False
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        for run in ([(ln, bold_first and i == 0)] if isinstance(ln, str) else ln):
            r = p.add_run()
            r.text = run[0]
            r.font.name, r.font.size, r.font.bold = "Calibri", Pt(size), run[1]
            r.font.color.rgb = run[2] if len(run) > 2 else color
    return b


def arrow(sh, x1, y1, x2, y2, color=INK, width=1.5, dash=False, head=True):
    c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle"); t.set("w", "med"); t.set("len", "med")
    return c


def panel(sh, x, y, w, h, fill=PANEL):
    p = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    p.fill.solid(); p.fill.fore_color.rgb = fill
    p.line.color.rgb = EDGE; p.line.width = Pt(1)
    p.shadow.inherit = False
    p.adjustments[0] = 0.06
    return p


# -------------------------------------------------- the U layout (slides 1-2)
BW, BH = 2.5, 0.62
ROWS = [2.05, 2.92, 3.79, 4.66]
BN_Y = 5.53
ENC_X = [0.75 + 0.62 * i for i in range(4)]
DEC_X = [13.33 - 0.75 - BW - 0.62 * i for i in range(4)]
BN_X = 13.33 / 2 - BW / 2
SIZES = ["8 × 512 × 512", "16 × 256 × 256", "32 × 128 × 128", "64 × 64 × 64"]


def draw_u(s, enc_ops, bn_ops, dec_ops, enc_fill, dec_fill, bn_fill, top_left, top_right):
    g = s.shapes.add_group_shape()
    sh = g.shapes
    text(sh, 0.75, 1.68, 6.0, 0.32, [top_left], size=11, color=GREY)
    text(sh, 6.6, 1.68, 6.0, 0.32, [top_right], size=11, color=GREY, align=PP_ALIGN.RIGHT)
    for i in range(4):
        box(sh, ENC_X[i], ROWS[i], BW, BH, enc_fill,
            [f"Encoder {i + 1} · {SIZES[i]}", enc_ops[i]], size=10.5)
        box(sh, DEC_X[i], ROWS[i], BW, BH, dec_fill,
            [f"Decoder {i + 1} · {SIZES[i]}", dec_ops[i]], size=10.5)
        # skip connection, across the U
        arrow(sh, ENC_X[i] + BW, ROWS[i] + BH / 2, DEC_X[i], ROWS[i] + BH / 2,
              color=GREY, width=1.25, dash=True)
    text(sh, ENC_X[0] + BW, ROWS[0] + BH / 2 - 0.27, DEC_X[0] - ENC_X[0] - BW, 0.25,
         ["skip connection: copy the encoder's features and join them to the decoder's"],
         size=10, color=GREY, align=PP_ALIGN.CENTER)
    for i in range(3):            # down the left side, up the right side
        xa = ENC_X[i + 1] + 0.3
        arrow(sh, xa, ROWS[i] + BH, xa, ROWS[i + 1])
        text(sh, xa - 1.25, ROWS[i] + BH, 1.2, 0.25, ["max-pool: half size"], size=9,
             color=GREY, align=PP_ALIGN.RIGHT)
        xb = DEC_X[i] + BW - 0.3 - 0.62 + 0.62
        arrow(sh, DEC_X[i + 1] + BW - 0.3, ROWS[i + 1], DEC_X[i + 1] + BW - 0.3, ROWS[i] + BH)
        text(sh, DEC_X[i + 1] + BW - 0.25, ROWS[i] + BH, 1.25, 0.25, ["up-conv: double size"],
             size=9, color=GREY)
    box(sh, BN_X, BN_Y, BW, BH, bn_fill, ["Bottleneck · 128 × 32 × 32", bn_ops], size=10.5)
    arrow(sh, ENC_X[3] + BW - 0.15, ROWS[3] + BH, BN_X, BN_Y + BH / 2)
    arrow(sh, BN_X + BW, BN_Y + BH / 2, DEC_X[3] + 0.15, ROWS[3] + BH)
    return g


# ======================================================================= 1
s = new_slide("Starting point: a basic U-Net")
draw_u(s,
       enc_ops=["2 × (3×3 conv → ReLU)"] * 4,
       bn_ops="2 × (3×3 conv → ReLU)",
       dec_ops=["join skip → 2 × (3×3 conv → ReLU)"] * 4,
       enc_fill=MIDNAVY, dec_fill=MIDNAVY, bn_fill=NAVY,
       top_left="▼ Input: spectrogram, 1 × 512 × 512 (time × frequency)",
       top_right="Output: an RFI probability for every pixel ▲")
panel(s.shapes, 0.75, 5.45, 3.70, 1.15)
text(s.shapes, 0.9, 5.52, 3.45, 1.05, [
    [("Encoder (left): ", True), ("shrinks the image step by step and learns WHAT is in it.", False)],
    [("Decoder (right): ", True), ("grows it back to full size and learns WHERE it is.", False)],
], size=11, space_after=6)
panel(s.shapes, 8.90, 5.45, 3.70, 1.15)
text(s.shapes, 9.05, 5.52, 3.45, 1.05, [
    [("Skip connections: ", True), ("pass fine detail straight across, so thin RFI lines are not lost "
                                    "when the image is shrunk.", False)],
    [("Numbers: ", True), ("channels × height × width.", False)],
], size=11, space_after=6)
panel(s.shapes, 0.75, 6.72, 11.85, 0.42, fill=PALE_RED)
text(s.shapes, 0.88, 6.75, 11.6, 0.36, [
    [("Drawn at our model's sizes so slides 1 and 2 line up box for box. ", True, RED),
     ("Akeret's tf_unet (our baseline) is this design, but shrinks only twice (32 → 64 → 128 channels, "
      "465,986 parameters) and trims the border, so 512 × 512 in gives 472 × 472 out.", False)],
], size=10, anchor=MSO_ANCHOR.MIDDLE)

# ======================================================================= 2
s = new_slide(f"{NAME}: the same U, changed inside")
draw_u(s,
       enc_ops=["ResBlock ②④ → Strip ① → ECA ③"] * 3 + ["ResBlock ②④ → ECA ③"],
       bn_ops="ResBlock ②④ → Strip ①",
       dec_ops=["join skip → ResBlock ②④ → ECA ③"] * 4,
       enc_fill=NAVY, dec_fill=NAVY, bn_fill=MIDNAVY,
       top_left="▼ Input: spectrogram, 1 × 512 × 512 — no border trimming ⑦",
       top_right="1×1 conv → 2 raw scores → softmax ⑤ → RFI probability ▲")
panel(s.shapes, 0.75, 5.40, 3.70, 1.72, fill=PALE_RED)
text(s.shapes, 0.88, 5.45, 3.48, 1.65, [
    [("① Strip convolutions", True, RED), (" — 7/11/21-px windows (enc. 1–3, bottleneck)", False)],
    [("② Residual shortcut", True, RED), (" — inside every ResBlock", False)],
    [("③ ECA attention", True, RED), (" — after every encoder and decoder block", False)],
    [("④ GroupNorm", True, NAVY), (" — in every ResBlock and strip (tf_unet: none)", False)],
    [("⑤ Raw scores", True, NAVY), (" — no ReLU before softmax (tf_unet has one)", False)],
], size=10, space_after=2)
panel(s.shapes, 8.90, 5.40, 3.70, 1.72)
text(s.shapes, 9.03, 5.45, 3.48, 1.65, [
    [("⑥ Dice + cross-entropy loss", True, NAVY), (" (tf_unet: cross-entropy only)", False)],
    [("⑦ Same padding", True, NAVY), (" — every pixel gets a prediction", False)],
    [("⑧ No class weighting", True, NAVY), (" — tf_unet has none either; our earlier version used it and scored lower", False)],
    [("⑨ Size", True, NAVY), (": shrinks 4 times, 8 → 128 channels, 593,842 parameters", False)],
    [("Red = the three added modules ①–③", True, RED)],
], size=10, space_after=2)

# ======================================================================= 3
s = new_slide("Inside one level, step by step")
X0 = 2.95            # where the flow diagrams start
ROW = [1.82, 3.40, 5.48]
heads = [("ResBlock  ②④", "Two 3×3 convolutions, each tidied by GroupNorm; the input is added back at the end. Dropout works in training only."),
         ("Strip module  ①", "Looks at wider areas (7, 11, 21 px) cheaply: a thin filter along frequency, then one along time."),
         ("ECA attention  ③", "Gives each feature map a weight between 0 and 1, then scales it.")]
heights = [1.45, 1.95, 1.42]
for (hd, sub), y, hh in zip(heads, ROW, heights):
    panel(s.shapes, 0.75, y, 11.85, hh)
    text(s.shapes, 0.88, y + 0.1, 1.95, hh - 0.2, [[(hd, True, RED)], sub], size=11,
         anchor=MSO_ANCHOR.MIDDLE, space_after=4)

# --- row A: ResBlock
g = s.shapes.add_group_shape(); sh = g.shapes
y = ROW[0] + 0.74; h = 0.5; gap = 0.15
steps = [("input", 0.58, LIGHT, INK), ("3×3 conv", 0.82, MIDNAVY, WHITE), ("GroupNorm ④", 1.15, NAVY, WHITE),
         ("ReLU", 0.6, MIDNAVY, WHITE), ("Dropout", 0.8, GREY, WHITE), ("3×3 conv", 0.82, MIDNAVY, WHITE),
         ("GroupNorm ④", 1.15, NAVY, WHITE)]
x = X0; xs = []
for lab, w, f, c in steps:
    box(sh, x, y, w, h, f, [lab], size=10, color=c); xs.append((x, w)); x += w
    arrow(sh, x, y + h / 2, x + gap, y + h / 2); x += gap
plus = box(sh, x, y + 0.07, 0.36, 0.36, RED, ["+"], size=14, shape=MSO_SHAPE.OVAL)
xp = x; x += 0.36
arrow(sh, x, y + h / 2, x + gap, y + h / 2); x += gap
box(sh, x, y, 0.6, h, MIDNAVY, ["ReLU"], size=10); x += 0.6
arrow(sh, x, y + h / 2, x + gap, y + h / 2); x += gap
box(sh, x, y, 0.68, h, LIGHT, ["output"], size=10, color=INK)
xin = xs[0][0] + xs[0][1] / 2; top = y - 0.36; xm = xp + 0.18
arrow(sh, xin, y, xin, top, color=RED, width=2, head=False)
sx = xin + 0.55; sw = 0.8
arrow(sh, xin, top, sx, top, color=RED, width=2, head=False)
box(sh, sx, top - 0.14, sw, 0.28, WHITE, [[("1×1 conv", True, RED)]], size=9.5, line=RED)
arrow(sh, sx + sw, top, xm, top, color=RED, width=2, head=False)
arrow(sh, xm, top, xm, y + 0.07, color=RED, width=2)
text(sh, sx + sw + 0.12, top - 0.27, 7.8, 0.22,
     [[("② shortcut: ", True, RED), ("input goes around the block and is added back (via a 1×1 conv: "
                                     "every ResBlock here changes the channel count)", False, RED)]],
     size=9)

# --- row B: Strip module (4 parallel paths)
g = s.shapes.add_group_shape(); sh = g.shapes
yc = ROW[1] + heights[1] / 2
box(sh, X0, yc - 0.25, 0.62, 0.5, LIGHT, ["input"], size=10.5, color=INK)
paths = ["copy (unchanged)", "1×7  →  7×1", "1×11  →  11×1", "1×21  →  21×1"]
px = X0 + 1.05; pw = 1.7; ph = 0.3
ys = [ROW[1] + 0.12 + k * 0.38 for k in range(4)]
for k, (lab, yy) in enumerate(zip(paths, ys)):
    box(sh, px, yy, pw, ph, LIGHT if k == 0 else RED, [lab], size=10,
        color=INK if k == 0 else WHITE, bold_first=False)
    arrow(sh, X0 + 0.62, yc, px, yy + ph / 2, color=GREY, width=1.1)
cx = px + pw + 0.45
box(sh, cx, yc - 0.3, 1.15, 0.6, MIDNAVY, ["join all 4", "(concatenate)"], size=10)
for yy in ys:
    arrow(sh, px + pw, yy + ph / 2, cx, yc, color=GREY, width=1.1)
x = cx + 1.15
for lab, w in [("1×1 conv (mix)", 1.15), ("GroupNorm ④", 1.2), ("ReLU", 0.7)]:
    arrow(sh, x, yc, x + gap, yc); x += gap
    box(sh, x, yc - 0.25, w, 0.5, NAVY if "Group" in lab else MIDNAVY, [lab], size=10.5); x += w
arrow(sh, x, yc, x + gap, yc); x += gap
box(sh, x, yc - 0.25, 0.75, 0.5, LIGHT, ["output"], size=10.5, color=INK)
text(sh, px, ROW[1] + heights[1] - 0.3, 7.5, 0.22,
     ["Each path = 1×K along frequency, then K×1 along time, each channel on its own → together it covers a K×K square (measured)"],
     size=9, color=GREY)

# --- row C: ECA
g = s.shapes.add_group_shape(); sh = g.shapes
y = ROW[2] + 0.38; h = 0.55
x = X0
box(sh, x, y, 0.62, h, LIGHT, ["input"], size=10.5, color=INK); xin = x + 0.31; x += 0.62
for lab, w, f in [("average each feature map → 1 number", 1.6, MIDNAVY),
                  ("1-D conv (k = 3) across neighbouring channels", 1.55, MIDNAVY),
                  ("sigmoid → weight 0 to 1", 1.3, RED)]:
    arrow(sh, x, y + h / 2, x + gap, y + h / 2); x += gap
    box(sh, x, y, w, h, f, [lab], size=10, bold_first=False); x += w
arrow(sh, x, y + h / 2, x + gap, y + h / 2); x += gap
mult = box(sh, x, y + 0.08, 0.4, 0.4, RED, ["×"], size=14, shape=MSO_SHAPE.OVAL); xm = x + 0.2; x += 0.4
arrow(sh, x, y + h / 2, x + gap, y + h / 2); x += gap
box(sh, x, y, 0.75, h, LIGHT, ["output"], size=10.5, color=INK)
bot = y + h + 0.22
arrow(sh, xin, y + h, xin, bot, color=GREY, width=1.25, head=False)
arrow(sh, xin, bot, xm, bot, color=GREY, width=1.25, head=False)
arrow(sh, xm, bot, xm, y + 0.48, color=GREY, width=1.25)
text(sh, xin + 0.15, bot - 0.02, 4.5, 0.22,
     ["the feature maps themselves go straight to × and are scaled by their weights"], size=9, color=GREY)

# ======================================================================= 4
s = new_slide("From output to RFI mask, and how it learns")
panel(s.shapes, 0.75, 1.76, 11.85, 1.28)
text(s.shapes, 0.88, 1.82, 2.0, 1.16, [[("Using the model", True, RED)], "after training"], size=12,
     anchor=MSO_ANCHOR.MIDDLE)
g = s.shapes.add_group_shape(); sh = g.shapes
y = 2.10; h = 0.6; x = 2.95; gap = 0.2
flow = [("spectrogram", 1.15, LIGHT, INK), ("scale to 0–1 (fixed range)", 1.3, MIDNAVY, WHITE),
        (f"{NAME}", 1.05, NAVY, WHITE), ("probability 0–1 for each pixel", 1.6, MIDNAVY, WHITE),
        ("threshold (picked on validation images)", 1.75, RED, WHITE), ("RFI mask: yes / no per pixel", 1.55, LIGHT, INK)]
for i, (lab, w, f, c) in enumerate(flow):
    box(sh, x, y, w, h, f, [lab], size=10, color=c, bold_first=False); x += w
    if i < len(flow) - 1:
        arrow(sh, x, y + h / 2, x + gap, y + h / 2); x += gap

panel(s.shapes, 0.75, 3.18, 11.85, 1.62)
text(s.shapes, 0.88, 3.24, 2.0, 1.5, [[("Training", True, RED)], "28,000 steps, one image per step"],
     size=12, anchor=MSO_ANCHOR.MIDDLE)
g = s.shapes.add_group_shape(); sh = g.shapes
y = 3.40; x = 2.95
box(sh, x, y, 1.7, 0.5, MIDNAVY, ["model's probabilities"], size=10, bold_first=False)
box(sh, x, y + 0.7, 1.7, 0.5, LIGHT, ["AOFlagger's RFI mask"], size=10, color=INK, bold_first=False)
lx = x + 1.7 + 0.35
box(sh, lx, y + 0.1, 2.3, 1.0, RED, ["loss = cross-entropy + Dice ⑥",
                                     "how wrong the model is; no class weighting ⑧"], size=10)
arrow(sh, x + 1.7, y + 0.25, lx, y + 0.45)
arrow(sh, x + 1.7, y + 0.95, lx, y + 0.75)
ux = lx + 2.3 + 0.35
box(sh, ux, y + 0.1, 2.25, 1.0, NAVY, ["Adam optimiser",
                                       "nudges all 593,842 numbers to lower the loss"], size=10)
arrow(sh, lx + 2.3, y + 0.6, ux, y + 0.6)
kx = ux + 2.25 + 0.35
box(sh, kx, y + 0.1, 1.95, 1.0, MIDNAVY, ["Every 1,400 steps",
                                          "check on 150 validation images; keep the best version"], size=10)
arrow(sh, ux + 2.25, y + 0.6, kx, y + 0.6)

panel(s.shapes, 0.75, 4.94, 11.85, 1.80, fill=PALE_RED)
text(s.shapes, 0.95, 5.04, 11.5, 1.64, [
    [("How it is scored: ", True, RED),
     ("109 test images labelled by a human expert; F1 at the best threshold (the standard for this "
      "benchmark); 3 runs each, mean ± spread.", False)],
    [("Our model 0.6603 ± 0.0040", True), ("  vs  Akeret's tf_unet 0.5482 ± 0.0139  →  about +0.11.", False)],
    [("Which parts matter? ", True, RED),
     ("Removing ①②③ gives 0.6585 — not significant (p = 0.64). "
      "Removing ④ as well gives 0.6495 — also not significant (p = 0.48).", False)],
    [("So the gain over tf_unet sits mainly in ⑤ ⑥ ⑦ ⑨ ", True),
     ("(raw output, Dice loss, same padding, size and training setup) — not yet tested one at a time.", False)],
], size=13, space_after=6)

prs.save(OUT)
print("saved", OUT, "|", len(prs.slides._sldIdLst), "slides")
