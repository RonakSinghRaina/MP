"""Lab-meeting deck: topic 11 (expert labels / active learning) on top of the hybrid.

Built on Weekly_Report_Week2_DiL_Lab_v2.pptx so the slides inherit the lab's
branding, exactly like build_hybrid_deck.py (whose helpers are repeated here).
Every chart is drawn from plain shapes and every table is a native table, so
both stay editable and resizable after importing into Google Slides.

All numbers are from RFI-project-context.md PART 27, printed by
analysis/active_learning/summarise_al.py (5-fold cross-validation over the 109
expert-labelled LOFAR images, 3 seeds).

    ~/torch-env/bin/python presentations/build_topic11_deck.py
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(HERE, ".."))
TEMPLATE = os.path.join(ROOT, "Weekly_Report_Week2_DiL_Lab_v2.pptx")
OUT = os.path.join(ROOT, "Weekly_Report_Topic11_DiL_Lab.pptx")
DATE = "11/10/2026"

RED = RGBColor(0xE3, 0x18, 0x37)
INK = RGBColor(0x1E, 0x1E, 0x1E)
NAVY = RGBColor(0x0F, 0x20, 0x43)
GREY = RGBColor(0x5A, 0x61, 0x70)
LIGHT = RGBColor(0xB9, 0xBE, 0xC9)
MID = RGBColor(0x8E, 0x9A, 0xAF)
PANEL = RGBColor(0xF6, 0xF7, 0xF9)
EDGE = RGBColor(0xD9, 0xDD, 0xE4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE_RED = RGBColor(0xFB, 0xE3, 0xE7)

prs = Presentation(TEMPLATE)
for sld in list(prs.slides._sldIdLst):        # keep masters/layouts, drop old slides
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)
LAY_TITLE, LAY_OBJ = prs.slide_layouts[0], prs.slide_layouts[1]


# ----------------------------------------------------------------- helpers (as build_hybrid_deck.py)
def new_slide(title):
    s = prs.slides.add_slide(LAY_OBJ)
    tf = s.placeholders[0].text_frame
    tf.clear()
    r = tf.paragraphs[0].add_run()
    r.text = "➢ " + title
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Calibri", Pt(26), True, RED
    for ph in list(s.placeholders):           # drop the layout's empty footer boxes
        if ph.placeholder_format.idx != 0 and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)
    return s


def text(shapes, x, y, w, h, paras, size=15, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6):
    """paras: list of paragraphs; each paragraph is a str or list of
    (text, bold[, color]) runs."""
    box = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = [(para, False)] if isinstance(para, str) else para
        for run in runs:
            t, b = run[0], run[1]
            c = run[2] if len(run) > 2 else color
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Calibri", Pt(size), b, c
    return box


def rect(shapes, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.0):
    r = shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line
        r.line.width = Pt(lw)
    r.shadow.inherit = False
    r.text_frame.text = ""
    return r


def label_in(shape, t, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = t
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Calibri", Pt(size), bold, color


def line(shapes, x1, y1, x2, y2, color=GREY, width=1.25):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                             Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def panel(shapes, x, y, w, h, fill=PANEL):
    return rect(shapes, x, y, w, h, fill, line=EDGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def bar_chart(slide, x, y, w, h, labels, values, colors, vmax, errors=None,
              fmt="{:.3f}", label_size=12, value_size=14):
    """Vertical bar chart made of plain shapes, grouped as one object.
    Bars start at zero so their heights are honest."""
    g = slide.shapes.add_group_shape()
    sh = g.shapes
    n = len(values)
    label_h = 0.62
    plot_h = h - label_h - 0.35           # room for value labels above the bars
    base_y = y + 0.35 + plot_h
    slot = w / n
    bw = slot * 0.58
    line(sh, x, base_y, x + w, base_y, color=GREY, width=1.25)
    for i, (lab, v, col) in enumerate(zip(labels, values, colors)):
        cx = x + slot * (i + 0.5)
        bh = plot_h * v / vmax
        rect(sh, cx - bw / 2, base_y - bh, bw, bh, col)
        top = base_y - bh
        if errors and errors[i]:
            eh = plot_h * errors[i] / vmax
            line(sh, cx, top - eh, cx, top + eh, color=INK, width=1.25)
            line(sh, cx - 0.08, top - eh, cx + 0.08, top - eh, color=INK, width=1.25)
            line(sh, cx - 0.08, top + eh, cx + 0.08, top + eh, color=INK, width=1.25)
            top -= eh
        vtxt = fmt.format(v) + (f" ± {errors[i]:.3f}" if errors and errors[i] else "")
        text(sh, cx - slot / 2, top - 0.36, slot, 0.32, [[(vtxt, True, col if col in (RED,) else INK)]],
             size=value_size, align=PP_ALIGN.CENTER, space_after=0)
        text(sh, cx - slot / 2, base_y + 0.06, slot, label_h, lab.split("\n"), size=label_size,
             color=INK, align=PP_ALIGN.CENTER, space_after=0)
    return g


def table(slide, rows, x, y, w, col_w, row_h=0.42, size=13):
    nr, nc = len(rows), len(rows[0])
    t = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                               Inches(row_h * nr)).table
    t.first_row = False
    t.horz_banding = False
    tot = sum(col_w)
    for i, cw in enumerate(col_w):
        t.columns[i].width = Emu(int(Inches(w) * cw / tot))
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if ri == 0 else (WHITE if ri % 2 else RGBColor(0xF3, 0xF5, 0xF8))
            c.margin_left = c.margin_right = Inches(0.08)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            c.text_frame.word_wrap = True
            p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.name, r.font.size = "Calibri", Pt(size)
            r.font.bold = ri == 0 or ci == 0
            r.font.color.rgb = WHITE if ri == 0 else (NAVY if ci == 0 else INK)
    return t


# ======================================================================= 1
s = prs.slides.add_slide(LAY_TITLE)
ph = s.placeholders[0]
ph.left, ph.top, ph.width, ph.height = Inches(1.78), Inches(4.86), Inches(10.0), Inches(0.7)
tf = ph.text_frame
tf.clear()
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Weekly Report – Topic 11: Expert Labels"
r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Calibri", Pt(36), True, RED
text(s.shapes, 1.78, 5.62, 10.0, 0.45, [DATE], size=21, align=PP_ALIGN.CENTER)
text(s.shapes, 0.08, 6.62, 4.4, 0.85,
     [[("Ronak Singh (BS-MS 4th Year)", True)], [("IMS23319", True)], [("School of Physics", True)]],
     size=14, space_after=0)

# ======================================================================= 2
s = new_slide("Topic 11: expert labels on our hybrid")
panel(s.shapes, 0.70, 1.75, 5.75, 4.15)
text(s.shapes, 0.92, 1.88, 5.35, 3.95, [
    [("What we did", True, RED)],
    [("•  The problem: ", True), ("our hybrid learned from AOFlagger's labels, which agree "
                                  "with the expert only at F1 0.570.", False)],
    [("•  Expert labels exist for only 109 images", True), (" — the test images.", False)],
    [("•  So we split the 109 into 5 parts. ", True), ("Each part is tested by a model that "
                                                       "learned from the other parts' expert labels "
                                                       "(68–69 images; 19 more kept for checking).", False)],
    [("•  Same hybrid, no changes: ", True), ("only a short extra training (600 steps).", False)],
], size=16, space_after=10)
bar_chart(s, 6.85, 1.80, 5.75, 4.10,
          ["No extra\ntraining", "+ expert labels\n(600 steps)", "+ expert labels\n(trained longer)"],
          [0.6603, 0.6747, 0.6676], [NAVY, RED, MID], vmax=0.75,
          errors=[0.0040, 0.0083, 0.0157], label_size=13, value_size=14)
text(s.shapes, 0.70, 6.02, 11.9, 0.85, [
    [("+0.014 F1 (higher for all 3 seeds, p = 0.03): 0.660 → 0.675 on all 109 images. ", True, RED),
     ("Level with RFDL's best result with expert labels (0.675), not above it.", True, NAVY)],
    [("max F1, mean ± spread of 3 runs. Results with expert labels must be compared with other expert-label "
      "results (RFDL: 0.669 / 0.675), not with the 0.660 table.", False)],
], size=13, color=GREY, space_after=2)

# ======================================================================= 3
s = new_slide("Which patches should the expert label?")
t = table(s, [
    ["Patches labelled (600 steps)", "5%", "10%", "20%"],
    ["Where the model is most uncertain", "0.6527", "0.6551", "0.6604"],
    ["Where the model sees most RFI", "0.6572", "0.6514", "0.6630"],
    ["At random", "0.6667", "0.6664", "0.6685"],
    ["Whole images at random", "0.6654", "0.6568", "0.6629"],
], x=0.70, y=1.85, w=7.05, col_w=[3.6, 1.15, 1.15, 1.15], row_h=0.52, size=15)
for ci in range(4):                                   # highlight the winning row
    c = t.cell(3, ci)
    c.fill.solid()
    c.fill.fore_color.rgb = PALE_RED
text(s.shapes, 0.70, 4.55, 7.05, 0.9, [
    "max F1 on all 109 images, mean of 3 runs. No extra training: 0.660 · all expert labels: 0.675.",
    "Share of the label pool: 5% ≈ 3 images, 20% ≈ 14 images.",
], size=12, color=GREY, space_after=2)
panel(s.shapes, 8.05, 1.75, 4.55, 3.95)
text(s.shapes, 8.25, 1.88, 4.15, 3.80, [
    [("The \"smart\" choices lose", True, RED)],
    [("Uncertain and most-RFI patches did worse than random ones in ", False),
     ("17 of 18", True), (" tests.", False)],
    [("Why: ", True), ("they are full of RFI — 3.2–7.4% of their pixels, against 0.76% overall. "
                       "The model learns \"RFI is common\" and flags too much.", False)],
    [("Train longer?", True, RED)],
    [("No: worse in ", False), ("12 of 13", True), (" settings. 19 checking images reward lucky "
                                                    "checkpoints.", False)],
], size=16, space_after=9)
text(s.shapes, 0.70, 6.10, 11.9, 0.6, [
    [("Rule of thumb: ", True, NAVY), ("train about 600 steps; if the expert's time is short, label random "
                                       "patches.", False)],
], size=16)

# ======================================================================= 4
s = new_slide("To Do Next")
text(s.shapes, 1.1, 1.85, 11.0, 4.8, [
    [("•  Dice-off run (3 seeds, ~2 h): ", True),
     ("is the Dice loss what puts our hybrid ahead of tf_unet?", False)],
    [("•  A possible fix for patch choice: ", True),
     ("keep AOFlagger's labels outside the chosen patches, so the model still sees normal "
      "background (not tested yet).", False)],
    [("•  Report: ", True), ("new section 5.6 with these results.", False)],
], size=19, space_after=16)

prs.save(OUT)
print("saved", OUT, "|", len(prs.slides._sldIdLst), "slides")
